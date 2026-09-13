"""Modern local workbench for the novel production system."""

from __future__ import annotations

import time
from datetime import datetime
from pathlib import Path
import json
import html
import re
from typing import Any

import streamlit as st

from core.analysis_skill import (
    analysis_to_markdown,
    build_analysis_skill_prompt,
    build_skill_zip,
    normalize_skill_markdown,
    skill_slug,
)
from core.job_manager import JobControl, JobManager, JobRecord
from core.models import AnalysisFile, LLMRequest, NovelMaterial
from core.skill_library import import_package, list_packages, package_context, packages_for_module
from core.studio_chat import (
    STUDIO_MODULES,
    build_studio_prompt,
    load_sessions,
    new_session,
    save_session,
)
from core.workspace import DEFAULT_WORKSPACE_ID, WorkspaceStore
from services.llm_client import LLMClient
from services.analyzer import Analyzer
from services.distiller import Distiller
from services.text_cleaner import TextCleaner
from services.quality_reviewer import review_chapter
from services.skill_evaluator import compare_skills


MODULES = {
    "home": ("⌂", "工作台"),
    "material": ("◈", "素材拆解"),
    "skill": ("◆", "Skill管理"),
    "creation": ("✎", "分步创作"),
    "chapters": ("▤", "章节写作"),
    "distill": ("✦", "规则蒸馏"),
    "tracking": ("◎", "设定追踪"),
    "feedback": ("↻", "迭代引擎"),
    "review": ("✓", "质量审查"),
    "settings": ("⚙", "API设置"),
}

WORKBENCH_COLUMNS = (1, 1)
FILE_WORKSPACE_COLUMNS = (2, 1)


def _css() -> None:
    st.markdown("""
    <style>
    :root { --ink:#252735; --muted:#73798a; --line:#e5e7ee; --soft:#f5f6fa; --panel:#fff; --purple:#625bd7; }
    [data-testid="stAppViewContainer"] { background:var(--soft); color:var(--ink); }
    [data-testid="stHeader"] { background:rgba(245,246,250,.96); }
    section[data-testid="stSidebar"] { background:#20222d; }
    section[data-testid="stSidebar"] * { color:#ececf3; }
    .wb-brand { display:flex; align-items:center; gap:10px; padding:8px 0 20px; font-weight:700; }
    .wb-mark { display:inline-grid; place-items:center; width:34px; height:34px; border-radius:9px; background:#6d64df; color:#fff; font-size:20px; }
    .wb-top { display:flex; justify-content:space-between; align-items:center; background:#fff; border:1px solid var(--line); border-radius:10px; padding:15px 20px; margin-bottom:15px; }
    .wb-title { font-size:24px; font-weight:700; }
    .wb-kicker,.wb-muted { color:var(--muted); font-size:12px; }
    .wb-pill { background:#e7f5f2; color:#13766e; border-radius:999px; padding:5px 10px; font-size:12px; }
    .wb-panel { background:#fff; border:1px solid var(--line); border-radius:10px; padding:17px; margin-bottom:14px; }
    .wb-panel h3 { margin:0 0 6px; font-size:17px; }
    .wb-stat { background:#fff; border:1px solid var(--line); border-radius:10px; padding:15px; min-height:88px; }
    .wb-stat-value { font-size:24px; font-weight:700; margin-top:8px; }
    .wb-file-folder { color:#5e56d1; font-weight:600; margin-top:8px; font-size:13px; }
    .wb-file { padding-left:15px; color:#555b6e; white-space:nowrap; overflow:hidden; text-overflow:ellipsis; font-size:12px; }
    .wb-file-workspace { display:flex; align-items:end; justify-content:space-between; margin:0 0 10px; padding:2px 2px 9px; border-bottom:1px solid var(--line); }
    .wb-file-workspace h3 { margin:0; color:var(--ink); font-size:17px; }
    .wb-file-empty { min-height:560px; display:grid; place-items:center; text-align:center; background:#f8fafc; border:1px dashed #b9c5d5; border-radius:8px; color:#65738a; padding:24px; }
    .wb-file-empty strong { display:block; color:#2b3b54; font-size:16px; margin-bottom:6px; }
    .wb-file-path { min-height:34px; padding:7px 10px; background:#eef2f7; border:1px solid #d3dae5; border-radius:6px; color:#33445c; font-size:12px; overflow:hidden; text-overflow:ellipsis; white-space:nowrap; }
    .wb-tree-meta { color:#69778d; font-size:11px; margin-bottom:8px; }
    .wb-file-editor textarea { line-height:1.8 !important; font-size:15px !important; }
    .wb-skill { border:1px solid #454958; border-radius:8px; padding:8px 10px; margin:6px 0; }
    .wb-skill strong { display:block; font-size:12px; }
    .wb-skill span { display:block; color:#a7aabb; font-size:10px; margin-top:3px; }
    .wb-chat { background:#fff; border:1px solid var(--line); border-radius:10px; padding:13px; margin-top:14px; }
    .wb-msg { background:#f1f1ff; border-radius:8px; padding:8px 10px; margin:6px 0; font-size:13px; }
    .wb-msg.user { background:#e8f5f2; }
    div[data-testid="stButton"] > button { border-radius:8px; border-color:var(--line); min-height:37px; }
    div[data-testid="stButton"] > button[kind="primary"] { background:var(--purple); border-color:var(--purple); }
    /* High-contrast desktop palette */
    :root { --ink:#172033; --muted:#526078; --line:#cbd3df; --soft:#edf1f6; --panel:#ffffff; --purple:#315ea8; }
    [data-testid="stAppViewContainer"] { background:#edf1f6; color:#172033; }
    section[data-testid="stSidebar"] { background:#202b3c; border-right:1px solid #34445b; }
    section[data-testid="stSidebar"] * { color:#edf4ff; }
    section[data-testid="stSidebar"] .stCaption, section[data-testid="stSidebar"] small { color:#b9c7dc !important; }
    section[data-testid="stSidebar"] div[data-testid="stButton"] > button { background:#2c3a50; border:1px solid #465972; color:#f5f8ff !important; text-align:left; font-weight:600; }
    section[data-testid="stSidebar"] div[data-testid="stButton"] > button:hover { background:#3b5272; border-color:#8aa4c9; }
    section[data-testid="stSidebar"] div[data-testid="stFileUploader"] { background:#2a374b; border:1px solid #465972; border-radius:8px; }
    section[data-testid="stSidebar"] .wb-skill { border-color:#526783; background:#263449; }
    section[data-testid="stSidebar"] .wb-skill strong { color:#ffffff; }
    section[data-testid="stSidebar"] .wb-skill span { color:#c3d0e3; }
    .wb-top, .wb-panel, .wb-stat, .wb-chat { border-color:#cbd3df; box-shadow:0 2px 8px rgba(27,42,65,.05); }
    .wb-kicker,.wb-muted { color:#526078; }
    .wb-file-folder { color:#244f8f; }
    .wb-file { color:#33445c; }
    div[data-testid="stButton"] > button { color:#20304a; background:#ffffff; border-color:#bfc9d7; }
    div[data-testid="stButton"] > button[kind="primary"] { background:#315ea8; border-color:#315ea8; color:#fff !important; }
    div[data-testid="stTextInput"] input, div[data-testid="stTextArea"] textarea { color:#172033; background:#fff; border-color:#bfc9d7; }    body, button, input, textarea, select { font-family: "Aptos", "Microsoft YaHei", sans-serif; }
    .wb-top { border-left:4px solid #c46a2b; }
    .wb-mark { background:#c46a2b; }
    .wb-panel h3 { letter-spacing:0; color:#172033; }
    .wb-panel > .wb-muted { max-width:76ch; line-height:1.6; }
    .wb-chat { border-top:3px solid #4c8b87; }
    section[data-testid="stSidebar"] div[data-testid="stButton"] > button[kind="primary"] { background:#c46a2b; border-color:#c46a2b; color:#fff !important; }
    section[data-testid="stSidebar"] div[data-testid="stButton"] > button[kind="secondary"] { color:#eaf1fb !important; }
    button:focus-visible, input:focus-visible, textarea:focus-visible, select:focus-visible { outline:3px solid rgba(196,106,43,.45) !important; outline-offset:2px; }
    @media (max-width: 900px) {
        .wb-title { font-size:20px; }
        .wb-top { padding:12px 14px; }
    }    /* Sidebar uploader: dark surface with readable labels */
    section[data-testid="stSidebar"] div[data-testid="stFileUploader"] { background:#263449; border:1px solid #526783; padding:8px; }
    section[data-testid="stSidebar"] div[data-testid="stFileUploaderDropzone"] { background:#2d3c52; border:1px dashed #7f94b2; }
    section[data-testid="stSidebar"] div[data-testid="stFileUploaderDropzone"] * { color:#eef4fc !important; }
    section[data-testid="stSidebar"] div[data-testid="stFileUploaderDropzone"] button { background:#c46a2b !important; border-color:#c46a2b !important; color:#fff !important; }
    section[data-testid="stSidebar"] div[data-testid="stFileUploaderDropzone"] small { color:#cbd8e8 !important; }
    section[data-testid="stSidebar"] div[data-testid="stFileUploaderDropzone"] svg { fill:#eef4fc; }    /* Streamlit 1.63 renders the uploader dropzone as a section, not a div. */
    section[data-testid="stSidebar"] section[data-testid="stFileUploaderDropzone"],
    section[data-testid="stSidebar"] div[data-testid="stFileUploaderDropzone"] {
        background:#263449 !important;
        border:1px dashed #8fa4c2 !important;
        border-radius:8px !important;
    }
    section[data-testid="stSidebar"] section[data-testid="stFileUploaderDropzone"] > div,
    section[data-testid="stSidebar"] section[data-testid="stFileUploaderDropzone"] > div > div {
        background:#263449 !important;
    }
    section[data-testid="stSidebar"] section[data-testid="stFileUploaderDropzone"] * {
        color:#f4f7fb !important;
    }
    section[data-testid="stSidebar"] section[data-testid="stFileUploaderDropzone"] small {
        color:#c9d7e8 !important;
    }
    section[data-testid="stSidebar"] section[data-testid="stFileUploaderDropzone"] button {
        background:#c46a2b !important;
        border:1px solid #c46a2b !important;
        color:#ffffff !important;
    }    /* Final uploader contrast override: black surface + white copy. */
    section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] {
        background-color:#0b111a !important;
        background:#0b111a !important;
        border:1px solid #5d7393 !important;
        color:#ffffff !important;
        opacity:1 !important;
    }
    section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] > div,
    section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] > div > div,
    section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] > div > div > div {
        background-color:#0b111a !important;
        background:#0b111a !important;
        color:#ffffff !important;
    }
    section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] p,
    section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] span,
    section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] small,
    section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] label {
        color:#ffffff !important;
        opacity:1 !important;
    }
    section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] button {
        background:#c46a2b !important;
        background-color:#c46a2b !important;
        border-color:#e18a52 !important;
        color:#ffffff !important;
        opacity:1 !important;
    }
    /* Streamlit 1.63 uploader internals: force a theme-independent surface. */
    section[data-testid="stSidebar"] [data-testid="stFileUploader"],
    section[data-testid="stSidebar"] [data-testid="stFileUploader"] > section,
    section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"],
    section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] > div {
        background:#000000 !important;
        background-color:#000000 !important;
        border-color:#6f829d !important;
    }
    section[data-testid="stSidebar"] [data-testid="stFileUploader"] p,
    section[data-testid="stSidebar"] [data-testid="stFileUploader"] small,
    section[data-testid="stSidebar"] [data-testid="stFileUploader"] span,
    section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzoneInstructions"],
    section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzoneInstructions"] * {
        color:#ffffff !important;
        -webkit-text-fill-color:#ffffff !important;
        opacity:1 !important;
    }
    section[data-testid="stSidebar"] [data-testid="stFileUploader"] button {
        background:#c46a2b !important;
        background-color:#c46a2b !important;
        border:1px solid #e8a06d !important;
        color:#ffffff !important;
        -webkit-text-fill-color:#ffffff !important;
        opacity:1 !important;
    }
    section[data-testid="stSidebar"] [data-testid="stFileUploader"] svg {
        color:#ffffff !important;
        fill:#ffffff !important;
    }
    .wb-studio-title { display:flex; align-items:center; justify-content:space-between; gap:12px; margin-bottom:8px; }
    .wb-studio-title strong { color:#172033; font-size:16px; }
    .wb-studio-title span { color:#65738a; font-size:11px; }
    .wb-studio-welcome { min-height:250px; display:grid; place-items:center; text-align:center; padding:30px; color:#66748a; }
    .wb-studio-welcome strong { display:block; color:#24344e; font-size:18px; margin-bottom:7px; }
    .wb-studio-welcome p { max-width:42ch; margin:0; line-height:1.7; }
    .wb-studio-history { color:#69778d; font-size:11px; margin:2px 0 9px; }
    [data-testid="stChatMessage"] { background:#f8fafc; border:1px solid #dde3ec; border-radius:8px; padding:9px 11px; }
    [data-testid="stChatMessage"]:has([data-testid="chatAvatarIcon-user"]) { background:#edf6f4; border-color:#cfe3df; }
    [data-testid="stChatInput"] { border-color:#bdc8d8; }
    [data-testid="stVerticalBlock"][height="460px"][class*="st-key-studio_chat_messages_"],
    [data-testid="stLayoutWrapper"][height="460px"]:has([class*="st-key-studio_chat_messages_"]) {
        height:clamp(140px, calc(100dvh - 620px), 460px) !important;
    }
    [class*="st-key-studio_chat_messages_"] .wb-studio-welcome { min-height:130px; padding:20px; }
    [data-testid="stLayoutWrapper"][height="460px"]:has([class*="st-key-studio_chat_messages_"]) {
        flex-basis:clamp(140px, calc(100dvh - 620px), 460px) !important;
        min-height:0;
    }
    </style>
    """, unsafe_allow_html=True)


CONVERSATION_MODES = ("material", "creation", "chapters", "distill")
NAVIGATION_MODULES = {
    "home": MODULES["home"],
    "conversation": ("✎", "创作对话"),
    **{key: value for key, value in MODULES.items() if key not in ("home", *CONVERSATION_MODES)},
}


def _unified_conversation() -> bool:
    return st.session_state.get("module") == "conversation"


def _label(module: str) -> str:
    if module == "conversation":
        return "创作对话"
    return MODULES.get(module, MODULES["home"])[1]


def _sidebar(store: WorkspaceStore) -> str:
    with st.sidebar:
        st.markdown('<div class="wb-brand"><span class="wb-mark">✦</span><span>OpenWrite<br><small style="color:#9fa3b8">网文创作工作台</small></span></div>', unsafe_allow_html=True)
        module = st.session_state.setdefault("module", "conversation")
        if module in CONVERSATION_MODES:
            st.session_state["conversation_mode"] = module
            module = "conversation"
            st.session_state["module"] = module
        for key, (icon, label) in NAVIGATION_MODULES.items():
            if st.button(f"{icon}  {label}", key=f"nav2_{key}", use_container_width=True, type="primary" if key == module else "secondary"): 
                st.session_state["module"] = key
                st.rerun()
            if key == "conversation" and module == "conversation":
                st.divider()
                mode = st.session_state.get("conversation_mode", st.session_state.get("conversation_last_mode", "material"))
                if mode not in CONVERSATION_MODES:
                    mode = "material"
                _studio_history(store, mode)
                st.divider()
        st.divider()
        st.caption("附件 Skill 库")
        upload = st.file_uploader("上传 Skill ZIP", type=["zip"], key="library_upload", label_visibility="collapsed")
        if upload and st.session_state.get("library_upload_name") != upload.name:
            import_package(upload.name, upload.getvalue())
            st.session_state["library_upload_name"] = upload.name
            st.toast("Skill 已加入本地库")
        for package in (list_packages() if module != "conversation" else []):
            st.write(f"{package.name} · {package.description}")
        st.caption("v2.0 · 本地工作区")
    return module


def _skill_picker(module: str) -> list[str]:
    packages = packages_for_module(module)
    ids = [p.package_id for p in packages]
    if not ids:
        st.info("当前模块暂无可用 Skill")
        return []
    default = st.session_state.setdefault(f"selected_{module}", ids[:1])
    value = st.multiselect("本模块使用的 Skill", ids, default=[x for x in default if x in ids], format_func=lambda x: next((p.name for p in packages if p.package_id == x), x), key=f"picker2_{module}")
    st.session_state[f"selected_{module}"] = value
    return value


def _document_editor(store: WorkspaceStore) -> None:
    relative_path = st.session_state.get("open_file")
    if not relative_path:
        st.markdown(
            "<div class='wb-file-empty'><div><strong>选择文件查看内容</strong>"
            "从右侧目录打开章节、设定或分析结果，正文会显示在这里。</div></div>",
            unsafe_allow_html=True,
        )
        return
    try:
        current = store.read_file(relative_path)
    except Exception as exc:
        st.error(f"无法打开文件：{exc}")
        st.session_state.pop("open_file", None)
        return

    st.markdown(
        f"<div class='wb-file-path'>{html.escape(relative_path)}</div>",
        unsafe_allow_html=True,
    )
    st.caption("可直接修改全文；保存时会自动生成历史版本。")
    st.markdown("<div class='wb-file-editor'>", unsafe_allow_html=True)
    edited = st.text_area(
        "文件内容",
        value=current,
        height=610,
        key=f"open_file_editor_{relative_path}",
        label_visibility="collapsed",
    )
    st.markdown("</div>", unsafe_allow_html=True)
    save_col, close_col, delete_col = st.columns([1.4, 1, 1])
    if save_col.button("保存", type="primary", key=f"save_open_file_{relative_path}", use_container_width=True):
        store.write_file(relative_path, edited)
        version_name = f"{Path(relative_path).stem}_{store.new_id()}.md"
        store.save_text("迭代记录", version_name, edited)
        st.toast("文件已保存，并生成版本记录")
    if close_col.button("关闭", key=f"close_open_file_{relative_path}", use_container_width=True):
        st.session_state.pop("open_file", None)
        st.rerun()
    if delete_col.button("删除", key=f"delete_open_file_{relative_path}", use_container_width=True):
        st.session_state["pending_delete_file"] = relative_path
        st.rerun()

    if st.session_state.get("pending_delete_file") == relative_path:
        st.warning(f"确定删除“{relative_path}”吗？此操作不能撤销。")
        confirm_col, cancel_col = st.columns(2)
        if confirm_col.button("确认删除", type="primary", key=f"confirm_delete_file_{relative_path}", use_container_width=True):
            store.delete_file(relative_path)
            st.session_state.pop("open_file", None)
            st.session_state.pop("pending_delete_file", None)
            st.toast("文件已删除")
            st.rerun()
        if cancel_col.button("取消", key=f"cancel_delete_file_{relative_path}", use_container_width=True):
            st.session_state.pop("pending_delete_file", None)
            st.rerun()


def _files(store: WorkspaceStore) -> None:
    tree = store.list_tree()
    st.markdown("#### 文件目录")
    st.markdown(
        f"<div class='wb-tree-meta'>{len(tree)} 个文件夹 · 点击文件打开</div>",
        unsafe_allow_html=True,
    )
    query = st.text_input(
        "搜索文件",
        placeholder="搜索文件...",
        key="workspace_file_search",
        label_visibility="collapsed",
    ).strip().lower()

    with st.expander("新建文件夹"):
        new_folder_name = st.text_input(
            "文件夹名称",
            placeholder="例如：第二卷",
            key="new_folder_name",
        )
        if st.button(
            "新建文件夹",
            key="create_workspace_folder",
            use_container_width=True,
        ):
            try:
                store.create_folder(new_folder_name)
                st.toast("文件夹已创建")
                st.rerun()
            except Exception as exc:
                st.error(str(exc))

    open_file = st.session_state.get("open_file", "")
    visible_folders = 0
    for folder_index, folder in enumerate(tree):
        files = [
            name for name in folder["files"]
            if not query or query in name.lower() or query in folder["name"].lower()
        ]
        if query and not files and query not in folder["name"].lower():
            continue
        visible_folders += 1
        expanded = bool(query) or open_file.startswith(f"{folder['name']}/")
        with st.expander(f"{folder['name']}  ({len(folder['files'])})", expanded=expanded):
            if st.button(
                "在此新建文件",
                key=f"start_new_file_{folder_index}",
                use_container_width=True,
            ):
                st.session_state["new_file_target"] = folder["name"]
                st.rerun()

            if st.session_state.get("new_file_target") == folder["name"]:
                new_file_name = st.text_input(
                    "文件名称",
                    placeholder="例如：人物小传（自动添加 .md）",
                    key=f"new_file_name_{folder_index}",
                )
                create_col, cancel_col = st.columns(2)
                if create_col.button(
                    "创建",
                    type="primary",
                    key=f"create_workspace_file_{folder_index}",
                    use_container_width=True,
                ):
                    try:
                        path = store.create_file(folder["name"], new_file_name)
                        st.session_state["open_file"] = path.relative_to(store.root).as_posix()
                        st.session_state.pop("new_file_target", None)
                        st.toast("文件已创建")
                        st.rerun()
                    except Exception as exc:
                        st.error(str(exc))
                if cancel_col.button(
                    "取消",
                    key=f"cancel_new_file_{folder_index}",
                    use_container_width=True,
                ):
                    st.session_state.pop("new_file_target", None)
                    st.rerun()

            if not files:
                st.caption("暂无文件")
            for file_index, name in enumerate(files):
                relative_path = f"{folder['name']}/{name}"
                if st.button(
                    name,
                    key=f"tree_file_{folder_index}_{file_index}",
                    type="primary" if relative_path == open_file else "secondary",
                    use_container_width=True,
                ):
                    st.session_state["open_file"] = relative_path
                    st.session_state.pop("pending_delete_file", None)
                    st.rerun()

            if not folder.get("system", True):
                if st.button(
                    "删除此文件夹",
                    key=f"delete_folder_{folder_index}",
                    use_container_width=True,
                ):
                    st.session_state["pending_delete_folder"] = folder["name"]
                    st.rerun()
                if st.session_state.get("pending_delete_folder") == folder["name"]:
                    st.warning("文件夹内的文件也会被删除。")
                    confirm_col, cancel_col = st.columns(2)
                    if confirm_col.button(
                        "确认删除",
                        type="primary",
                        key=f"confirm_delete_folder_{folder_index}",
                        use_container_width=True,
                    ):
                        store.delete_folder(folder["name"])
                        if open_file.startswith(f"{folder['name']}/"):
                            st.session_state.pop("open_file", None)
                        st.session_state.pop("pending_delete_folder", None)
                        st.toast("文件夹已删除")
                        st.rerun()
                    if cancel_col.button(
                        "取消",
                        key=f"cancel_delete_folder_{folder_index}",
                        use_container_width=True,
                    ):
                        st.session_state.pop("pending_delete_folder", None)
                        st.rerun()
    if visible_folders == 0:
        st.info("没有找到匹配的文件。")


def _file_workspace(store: WorkspaceStore) -> None:
    st.markdown(
        "<div class='wb-file-workspace'><div><h3>项目文件工作区</h3>"
        "<div class='wb-muted'>查看、编辑和管理本地创作文件</div></div></div>",
        unsafe_allow_html=True,
    )
    editor_col, tree_col = st.columns(FILE_WORKSPACE_COLUMNS, gap="small")
    with editor_col:
        _document_editor(store)
    with tree_col:
        _files(store)
def _header(module: str) -> None:
    st.markdown(f"<div class='wb-top'><div><div class='wb-kicker'>默认工作区 · 本地数据</div><div class='wb-title'>{_label(module)}</div></div><span class='wb-pill'>● 已连接本地服务</span></div>", unsafe_allow_html=True)


def _job_controls(store: WorkspaceStore, module: str) -> None:
    manager = JobManager(store)
    job = manager.active_job(module)
    if not job:
        return
    st.markdown(f"<div class='wb-panel'><h3>任务：{job.title}</h3><div class='wb-muted'>{job.message} · {job.progress}%</div></div>", unsafe_allow_html=True)
    st.progress(job.progress / 100, text=job.message)
    c1, c2, c3 = st.columns(3)
    if job.status == "paused":
        if c1.button("▶ 继续", key=f"resume2_{job.job_id}"): manager.resume(job.job_id); st.rerun()
    elif job.status in {"queued", "running"}:
        if c1.button("Ⅱ 暂停", key=f"pause2_{job.job_id}"): manager.pause(job.job_id); st.rerun()
    if job.status in {"queued", "running", "paused"} and c2.button("■ 终止", key=f"cancel2_{job.job_id}"): manager.cancel(job.job_id); st.rerun()
    if c3.button("↻ 刷新", key=f"refresh2_{job.job_id}"): st.rerun()


def _start_job(store: WorkspaceStore, module: str, title: str, skills: list[str], payload: dict[str, Any]) -> None:
    manager = JobManager(store)
    if manager.active_job(module):
        st.warning(f"“{_label(module)}”已有任务在处理中，请先暂停、终止或等待它完成。")
        return

    def process(control: JobControl, job: JobRecord) -> str:
        for index in range(20):
            if not control.checkpoint():
                return ""
            time.sleep(0.25)
            manager.update_progress(job.job_id, index * 4, message=f"正在处理第 {index}/20 阶段")

        if module == "material":
            source_path = store.root / payload["source"]
            if source_path.suffix.lower() == ".docx":
                from docx import Document
                raw_text = "\n".join(paragraph.text for paragraph in Document(source_path).paragraphs)
            else:
                raw_text = source_path.read_text(encoding="utf-8", errors="replace")
            material = NovelMaterial(filename=source_path.name, raw_text=raw_text, file_size=source_path.stat().st_size)
            cleaner = TextCleaner()
            cleaned = cleaner.clean(material)
            cleaned.chunks = cleaner.chunk_text(cleaned.cleaned_text)
            store.save_text("小说资料", f"{source_path.stem}_cleaned.md", cleaned.cleaned_text)
            client = LLMClient()
            has_api = bool(
                (client._primary and client._primary.api_key)
                or (client._backup and client._backup.api_key)
            )
            analyses = Analyzer(client, package_context(skills)).analyze(cleaned) if has_api else []
            output = {
                "job_id": job.job_id,
                "module": module,
                "skills": skills,
                "skill_context": package_context(skills),
                "source": payload["source"],
                "cleaned_chars": len(cleaned.cleaned_text),
                "chunk_count": len(cleaned.chunks),
                "analyses": [item.model_dump(mode="json") for item in analyses],
                "created_at": store.timestamp(),
            }
            path = store.save_json("拆解结果", f"{source_path.stem}_{job.job_id}.json", output)
            return str(path.relative_to(store.root))

        if module == "distill":
            analysis_files = _distill_analysis_files(store)
            if not analysis_files:
                raise RuntimeError("没有可用的拆解分析结果，请先完成素材拆解并确认分析产物已保存。")
            chapter_files = [store.root / "章节内容" / name for name in payload.get("chapter_files", [])]
            review_files = [store.root / "迭代记录" / name for name in payload.get("review_files", [])]
            tracker_files = [store.root / "创作结果" / name for name in payload.get("tracker_files", [])]
            latest = json.loads(analysis_files[0].read_text(encoding="utf-8"))
            analysis_items = [AnalysisFile.model_validate(item) for item in latest.get("analyses", [])]
            if not analysis_items:
                raise RuntimeError("最新拆解结果没有有效分析条目，请重新运行素材拆解。")
            rule_set = Distiller().distill(analysis_items, name=title)
            evidence = {
                "chapter_files": [
                    {"file": path.name, "content": path.read_text(encoding="utf-8", errors="replace")[:8000]}
                    for path in chapter_files if path.is_file()
                ],
                "review_files": [
                    {"file": path.name, "report": json.loads(path.read_text(encoding="utf-8"))}
                    for path in review_files if path.is_file()
                ],
                "tracker_files": [
                    {"file": path.name, "content": path.read_text(encoding="utf-8", errors="replace")[:6000]}
                    for path in tracker_files if path.is_file()
                ],
            }
            path = store.save_json("蒸馏结果", f"{title}_{job.job_id}.json", {
                "job_id": job.job_id,
                "skills": skills,
                "skill_context": package_context(skills),
                "source_analysis": analysis_files[0].name,
                "rule_set": rule_set.model_dump(mode="json"),
                "evidence": evidence,
                "created_at": store.timestamp(),
            })
            return str(path.relative_to(store.root))
        folder = {"feedback": "迭代记录"}.get(module, "任务记录")
        path = store.save_json(folder, f"{module}_{job.job_id}.json", {"job_id": job.job_id, "module": module, "skills": skills, "skill_context": package_context(skills), "payload": payload, "created_at": store.timestamp()})
        return str(path.relative_to(store.root))

    manager.start(module, title, skills, process)
    st.rerun()


def _material_tools(store: WorkspaceStore, skills: list[str]) -> None:
    upload = st.file_uploader("上传小说或章节", type=["txt", "md", "docx"], key="material_upload2")
    if upload: st.info(f"已选择：{upload.name} · {upload.size / 1024:.1f} KB")
    if st.button("▶ 开始拆解", type="primary", disabled=upload is None):
        source = store.save_bytes("小说资料", upload.name, upload.getvalue())
        _start_job(store, "material", f"拆解 {upload.name}", skills, {"source": str(source.relative_to(store.root))})
    _job_controls(store, "material")
    st.markdown("<div class='wb-panel'><h3>最近结果</h3>", unsafe_allow_html=True)
    for job in [x for x in JobManager(store).list_jobs() if x.module == "material"][:6]: st.write(f"{job.status.upper()} · {job.title} · {job.output_file or job.message}")
    st.markdown("</div>", unsafe_allow_html=True)


def _material(store: WorkspaceStore) -> None:
    skills = _skill_picker("material")
    context = ""
    source_name = "未选择素材"
    with st.expander("素材文件与后台拆解", expanded=False):
        _material_tools(store, skills)
        sources = sorted(
            [p for p in (store.root / "小说资料").iterdir() if p.is_file() and p.suffix.lower() in {".txt", ".md", ".docx"}],
            key=lambda p: p.stat().st_mtime,
            reverse=True,
        )
        if sources:
            source_name = st.selectbox("本次对话使用的素材", [p.name for p in sources], key="material_chat_source")
            source = next(p for p in sources if p.name == source_name)
            if source.suffix.lower() == ".docx":
                from docx import Document
                context = "\n".join(p.text for p in Document(source).paragraphs)
            else:
                context = source.read_text(encoding="utf-8", errors="replace")
    context = context[:18000]
    analysis_files = _distill_analysis_files(store)
    if analysis_files:
        context += "\n\n最近完成的拆解结果：\n" + analysis_to_markdown(
            json.loads(analysis_files[0].read_text(encoding="utf-8"))
        )[:12000]
    editor_key = "studio_editor_material"
    st.session_state.setdefault(editor_key, st.session_state.get("studio_latest_material", ""))
    _studio_chat(
        store, "material", skills,
        {"素材文件": source_name, "已完成拆解数": len(analysis_files), "素材说明": "上下文仅含有限片段，不能声称已经读完全部素材"},
        context[:30000], "studio_latest_material", editor_key,
        "和 AI 讨论如何拆解素材、修正结论或生成 Skill...",
        (
            ("拆解结构", "请结合提供的素材拆解故事结构、人物目标和冲突，引用依据并区分事实与推断。"),
            ("提炼方法", "请从当前素材与讨论中提炼可复用的写作方法，说明适用范围。"),
            ("生成 Skill", "请结合我对拆解结果的修改和创作要求，生成完整的 SKILL.md 草稿；如缺少我的要求，先询问。"),
        ),
    )
    with st.expander("对话拆解草稿与保存", expanded=False):
        draft = st.text_area("拆解草稿", height=380, key=editor_key)
        name = st.text_input("成果名称", value="我的拆解写作法", key="material_draft_name")
        left, right = st.columns(2)
        if left.button("保存拆解稿", key="save_material_chat", use_container_width=True, disabled=not draft.strip()):
            path = store.save_text("拆解结果", f"对话拆解_{store.new_id()}.md", draft)
            st.session_state["open_file"] = path.relative_to(store.root).as_posix()
            st.success("拆解稿已保存")
        if right.button("保存为新 Skill", key="save_material_chat_skill", use_container_width=True, disabled=not draft.strip()):
            text = normalize_skill_markdown(name, f"基于素材拆解的{name}，用于指导小说创作。", draft)
            version = store.new_id()
            slug = skill_slug(name)
            path = store.save_text("Skill", f"{slug}_{version}.md", text)
            import_package(f"{slug}-{version}.zip", build_skill_zip(slug, text))
            st.session_state["open_file"] = path.relative_to(store.root).as_posix()
            st.success("新 Skill 已加入本地库")
    with st.expander("已完成拆解：编辑结果与生成 Skill", expanded=False):
        _analysis_skill_builder(store, skills)


def _distill_analysis_files(store: WorkspaceStore) -> list[Path]:
    """Return the newest拆解 outputs that contain actual analysis entries."""
    candidates = sorted((store.root / "拆解结果").glob("*.json"), key=lambda item: item.stat().st_mtime, reverse=True)
    valid: list[Path] = []
    for path in candidates:
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError, TypeError):
            continue
        if isinstance(data.get("analyses"), list) and data["analyses"]:
            valid.append(path)
    return valid


def _analysis_skill_builder(store: WorkspaceStore, skills: list[str]) -> None:
    """Edit a completed analysis and turn it into a reusable local Skill."""
    analysis_files = _distill_analysis_files(store)
    st.markdown(
        "<div class='wb-panel'><h3>拆解结果编辑与 Skill 生成</h3>"
        "<div class='wb-muted'>修改 AI 拆解结论，补充你的创作要求，再生成可继续编辑的新 Skill。</div></div>",
        unsafe_allow_html=True,
    )
    if not analysis_files:
        st.info("拆解任务完成并产出有效结果后，这里会自动出现结果编辑器。")
        return

    selected_name = st.selectbox(
        "选择拆解结果",
        [path.name for path in analysis_files],
        key="analysis_result_select",
    )
    selected_path = next(path for path in analysis_files if path.name == selected_name)
    source_marker = str(selected_path.resolve())
    if st.session_state.get("analysis_result_source") != source_marker:
        data = json.loads(selected_path.read_text(encoding="utf-8"))
        st.session_state["analysis_result_source"] = source_marker
        st.session_state["analysis_result_editor"] = analysis_to_markdown(data)
        st.session_state["analysis_skill_name"] = f"{selected_path.stem}写作法"
        st.session_state["analysis_author_direction"] = ""
        st.session_state.pop("analysis_skill_draft", None)

    edited_result = st.text_area(
        "可编辑拆解结果",
        height=420,
        key="analysis_result_editor",
        help="可以删除不准确的结论、改写规则、补充你从原文中观察到的方法。",
    )
    save_edit_col, open_edit_col = st.columns(2)
    if save_edit_col.button("保存拆解编辑稿", key="save_analysis_edit", use_container_width=True):
        edit_path = store.save_text("拆解结果", f"{selected_path.stem}_编辑稿.md", edited_result)
        st.session_state["open_file"] = str(edit_path.relative_to(store.root)).replace("\\", "/")
        st.success(f"编辑稿已保存：{edit_path.name}")
    if open_edit_col.button("在右侧查看编辑稿", key="open_analysis_edit", use_container_width=True):
        edit_path = store.root / "拆解结果" / f"{selected_path.stem}_编辑稿.md"
        if edit_path.exists():
            st.session_state["open_file"] = str(edit_path.relative_to(store.root)).replace("\\", "/")
            st.rerun()
        else:
            st.warning("请先保存拆解编辑稿。")

    st.markdown("### 结合你的想法生成新 Skill")
    skill_name = st.text_input("新 Skill 名称", key="analysis_skill_name")
    author_direction = st.text_area(
        "你的创作要求与补充说明",
        placeholder="例如：强调普通人视角；节奏快但不强行反转；感情线必须服务主线……",
        height=130,
        key="analysis_author_direction",
    )
    can_generate = bool(skill_name.strip() and author_direction.strip() and edited_result.strip())
    if st.button(
        "生成新的 Skill",
        type="primary",
        key="generate_analysis_skill",
        disabled=not can_generate,
        use_container_width=True,
    ):
        prompt = build_analysis_skill_prompt(skill_name, author_direction, edited_result[:30000])
        try:
            with st.spinner("AI 正在融合拆解结果和你的创作要求..."):
                generated = _generate(prompt, skills)
            st.session_state["analysis_skill_draft"] = normalize_skill_markdown(
                skill_name,
                author_direction[:400],
                generated,
            )
            st.rerun()
        except Exception as exc:
            st.error(f"生成失败：{exc}")

    if "analysis_skill_draft" not in st.session_state:
        return
    skill_draft = st.text_area(
        "新 Skill 内容（可继续修改）",
        height=460,
        key="analysis_skill_draft",
    )
    save_skill_col, discard_col = st.columns(2)
    if save_skill_col.button(
        "保存并加入 Skill 库",
        type="primary",
        key="save_analysis_skill",
        use_container_width=True,
    ):
        try:
            normalized = normalize_skill_markdown(skill_name, author_direction[:400], skill_draft)
            version_id = store.new_id()
            slug = skill_slug(skill_name)
            skill_path = store.save_text("Skill", f"{slug}_{version_id}.md", normalized)
            package_path = import_package(
                f"{slug}-{version_id}.zip",
                build_skill_zip(slug, normalized),
            )
            store.save_json("迭代记录", f"analysis_skill_{version_id}.json", {
                "source_analysis": selected_path.name,
                "edited_analysis": f"{selected_path.stem}_编辑稿.md",
                "skill_name": skill_name,
                "author_direction": author_direction,
                "skill_file": skill_path.name,
                "package_file": package_path.name,
                "created_at": store.timestamp(),
            })
            st.session_state["open_file"] = str(skill_path.relative_to(store.root)).replace("\\", "/")
            st.success(f"新 Skill 已保存并加入选择列表：{skill_path.name}")
        except Exception as exc:
            st.error(f"保存失败：{exc}")

    def clear_skill_draft() -> None:
        st.session_state.pop("analysis_skill_draft", None)

    discard_col.button(
        "清空 Skill 草稿",
        key="discard_analysis_skill",
        use_container_width=True,
        on_click=clear_skill_draft,
    )


def _distill(store: WorkspaceStore) -> None:
    if not _unified_conversation():
        st.subheader("规则蒸馏")
    analysis_files = _distill_analysis_files(store)
    chapter_files = sorted((store.root / "章节内容").glob("第*.md"))
    review_files = sorted((store.root / "迭代记录").glob("review_*.json"), key=lambda item: item.stat().st_mtime, reverse=True)
    tracker_files = sorted((store.root / "创作结果").glob("*_tracker.md")) + [store.root / "创作结果" / "serial_status.md"]
    tracker_files = [path for path in tracker_files if path.is_file()]
    evidence_count = len(chapter_files) + len(review_files) + len(tracker_files)
    if not _unified_conversation():
        st.caption(f"证据：{len(analysis_files)} 份拆解、{len(chapter_files)} 章正文")

    skills = _skill_picker("distill")
    with st.expander("蒸馏设置", expanded=False):
        name = st.text_input(
            "规则集名称",
            value=st.session_state.get("distill_rule_name", "我的小说写作规则"),
            key="distill_rule_name_input",
        )
        st.session_state["distill_rule_name"] = name
        st.caption("AI 会自动读取可用的拆解、正文、设定追踪和审查证据。")

    draft_key = "studio_latest_distill"
    editor_key = "studio_editor_distill"
    if editor_key not in st.session_state:
        st.session_state[editor_key] = st.session_state.get(draft_key, "")
    evidence_context = _distill_context(
        analysis_files,
        chapter_files,
        review_files,
        tracker_files,
    )
    _studio_chat(
        store,
        "distill",
        skills,
        {
            "规则集名称": name,
            "拆解结果数": len(analysis_files),
            "章节数": len(chapter_files),
            "设定追踪数": len(tracker_files),
            "审查记录数": len(review_files),
            "蒸馏目标": "形成有依据、可执行、可验证、可回滚的写作规则",
        },
        evidence_context,
        draft_key,
        editor_key,
        "和 AI 讨论如何提炼与改进规则...",
        (
            ("提炼规则", "请从现有证据中提炼最值得保留的可复用写作规则，并逐条标注依据。"),
            ("寻找问题", "请比较 Skill 规则与正文表现，找出失效规则、冲突规则和缺失规则。"),
            ("生成新版", "请根据当前讨论生成一份结构完整、可直接使用的新版 Skill 草稿。"),
        ),
    )

    with st.expander("规则草稿与确认", expanded=bool(st.session_state.get(editor_key)) and not _unified_conversation()):
        draft = st.text_area(
            "规则草稿",
            height=480,
            key=editor_key,
            label_visibility="collapsed",
        )
        save_col, skill_col = st.columns(2)
        if save_col.button("保存蒸馏结果", type="primary", key="save_distill_draft", use_container_width=True):
            if not name.strip() or not draft.strip():
                st.warning("请先填写规则集名称并通过对话形成规则草稿。")
            else:
                safe_name = re.sub(r"[^\w\u4e00-\u9fff-]+", "-", name).strip("-") or "写作规则"
                path = store.save_text("蒸馏结果", f"{safe_name}_{store.new_id()}.md", draft)
                st.session_state["open_file"] = str(path.relative_to(store.root)).replace("\\", "/")
                st.success(f"蒸馏结果已保存：{path.name}")
        if skill_col.button("确认为新 Skill", key="confirm_distill_skill", use_container_width=True):
            if not name.strip() or not draft.strip():
                st.warning("规则草稿不能为空。")
            else:
                safe_name = re.sub(r"[^\w\u4e00-\u9fff-]+", "-", name).strip("-") or "写作Skill"
                skill_id = store.new_id()
                skill_text = (
                    f"# {name}\n\n"
                    f"来源 Skill：{', '.join(skills) or '无'}\n\n"
                    f"## 已确认规则\n\n{draft}\n"
                )
                path = store.save_text("Skill", f"{safe_name}_{skill_id}.md", skill_text)
                skill_text = normalize_skill_markdown(name, f"用于指导{name}的写作规则。", skill_text)
                store.save_text("Skill", path.name, skill_text)
                import_package(f"{skill_slug(name)}-{skill_id}.zip", build_skill_zip(skill_slug(name), skill_text))
                store.save_json(
                    "迭代记录",
                    f"distill_confirm_{skill_id}.json",
                    {
                        "name": name,
                        "source_skills": skills,
                        "evidence_count": evidence_count,
                        "output": str(path.relative_to(store.root)),
                        "created_at": store.timestamp(),
                    },
                )
                st.session_state["open_file"] = str(path.relative_to(store.root)).replace("\\", "/")
                st.success(f"新 Skill 已确认并保存：{path.name}")

    with st.expander("高级：后台全量蒸馏任务"):
        st.caption("需要长时间批量处理时使用，可暂停、恢复或终止。")
        if st.button("启动后台全量蒸馏", key="start_full_distill_job", use_container_width=True):
            if not name.strip():
                st.warning("请先填写规则集名称。")
            elif not analysis_files or not chapter_files:
                st.warning("后台全量蒸馏至少需要一份拆解结果和一章正文。")
            else:
                _start_job(
                    store,
                    "distill",
                    name.strip(),
                    skills,
                    {
                        "mode": "full",
                        "context_chars": len(package_context(skills)),
                        "analysis_file": analysis_files[0].name,
                        "chapter_files": [path.name for path in chapter_files],
                        "review_files": [path.name for path in review_files],
                        "tracker_files": [path.name for path in tracker_files],
                        "evidence_count": evidence_count,
                    },
                )
        _job_controls(store, "distill")

def _skill(store: WorkspaceStore) -> None:
    st.markdown("<div class='wb-panel'><h3>Skill管理</h3><div class='wb-muted'>固定附件、用户上传和融合 Skill 都在这里独立管理。</div></div>", unsafe_allow_html=True)
    packages = list_packages()
    for package in packages:
        with st.container(border=True):
            c1, c2, c3 = st.columns([3, 1, 1])
            c1.markdown(f"**{package.name}** · {package.description}\n\n`{package.file_count} files` · `{', '.join(package.modules)}`")
            if c2.button("查看", key=f"skill_view2_{package.package_id}"): st.code(package.read_entry()[:4000], language="markdown")
            if c3.button("保存到工作区", key=f"skill_copy2_{package.package_id}"): store.save_bytes("Skill", package.zip_path.name, package.zip_path.read_bytes()); st.toast("已保存")
    selected = st.multiselect("融合 Skill", [p.package_id for p in packages], format_func=lambda x: next((p.name for p in packages if p.package_id == x), x))
    name = st.text_input("融合名称", value="我的融合Skill")
    if st.button("◆ 生成融合 Skill", type="primary", disabled=len(selected) < 2):
        path = store.save_json("Skill", f"{name}.json", {"name": name, "source_skills": selected, "context": package_context(selected), "created_at": store.timestamp()})
        st.success(f"已保存：{path.name}")
    _skill_version_compare(store)


def _skill_text(store: WorkspaceStore, identifier: str, packages: list[Any]) -> tuple[str, str]:
    if identifier.startswith("package:"):
        package_id = identifier.removeprefix("package:")
        package = next(item for item in packages if item.package_id == package_id)
        return package.name, package.read_entry()
    path = store.root / "Skill" / identifier
    return path.name, path.read_text(encoding="utf-8", errors="replace")


def _skill_version_compare(store: WorkspaceStore) -> None:
    packages = list_packages()
    candidates = [f"package:{item.package_id}" for item in packages]
    candidates += [path.name for path in sorted((store.root / "Skill").glob("*.md"))]
    if len(candidates) < 2:
        return
    st.markdown("### Skill 版本对比 / A-B 评估")
    st.caption("先用透明的本地指标比较结构完整度，再决定是否进入小规模试用。")
    a, b = st.columns(2)
    with a:
        left = st.selectbox("版本 A", candidates, key="skill_ab_a")
    with b:
        right = st.selectbox("版本 B", candidates, index=min(1, len(candidates) - 1), key="skill_ab_b")
    if st.button("比较两个版本", key="compare_skill_versions", disabled=left == right):
        name_a, text_a = _skill_text(store, left, packages)
        name_b, text_b = _skill_text(store, right, packages)
        result = compare_skills(name_a, text_a, name_b, text_b)
        result["created_at"] = store.timestamp()
        store.save_json("迭代记录", f"skill_ab_{store.new_id()}.json", result)
        st.session_state["skill_ab_result"] = result
    if st.session_state.get("skill_ab_result"):
        result = st.session_state["skill_ab_result"]
        st.success(f"建议关注：{result['winner']}")
        st.json(result)


def _project_context(store: WorkspaceStore) -> str:
    profile = store.project_profile()
    parts = [f"项目设定：{json.dumps(profile, ensure_ascii=False)}"]
    for path in sorted((store.root / "创作结果").glob("*.md")):
        parts.append(f"{path.name}：\n{path.read_text(encoding='utf-8', errors='replace')[:5000]}")
    return "\n\n".join(parts)


def _generate(prompt: str, skills: list[str], context_text: str = "") -> str:
    context = package_context(skills)
    if context_text:
        context += f"\n\n{context_text}"
    result = LLMClient().chat(LLMRequest(prompt=f"{context}\n\n用户任务：{prompt}", max_tokens=5000), use_cache=False)
    if not result.success: raise RuntimeError(result.error_message)
    return result.content.strip()


def _studio_sessions(store: WorkspaceStore, module: str) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Return in-memory sessions and the active session for one module."""
    sessions_key = f"studio_sessions_{module}"
    active_key = f"studio_active_{module}"
    if sessions_key not in st.session_state:
        sessions = load_sessions(store, module)[:12]
        st.session_state[sessions_key] = sessions or [new_session(store, module)]
    sessions = st.session_state[sessions_key]
    active_id = st.session_state.get(active_key)
    active = next((item for item in sessions if item["id"] == active_id), None)
    if active is None:
        active = sessions[0]
        st.session_state[active_key] = active["id"]
    return sessions, active


def _studio_history(store: WorkspaceStore, module: str) -> None:
    sessions, active = _studio_sessions(store, module)
    st.markdown(f"##### {_label(module)}对话")
    if st.button("新建对话", key=f"studio_new_{module}", type="primary", use_container_width=True):
        created = new_session(store, module)
        save_session(store, created)
        sessions.insert(0, created)
        st.session_state[f"studio_active_{module}"] = created["id"]
        st.session_state[f"studio_restore_{module}"] = True
        st.rerun()
    for session in sessions[:12]:
        if st.button(
            str(session.get("title") or "新对话")[:20],
            key=f"studio_session_{module}_{session['id']}",
            type="primary" if session["id"] == active["id"] else "secondary",
            use_container_width=True,
        ):
            st.session_state[f"studio_active_{module}"] = session["id"]
            st.session_state[f"studio_restore_{module}"] = True
            st.rerun()


def _studio_chat(
    store: WorkspaceStore,
    module: str,
    skills: list[str],
    metadata: dict[str, Any],
    context_text: str,
    draft_key: str,
    editor_key: str,
    placeholder: str,
    quick_prompts: tuple[tuple[str, str], ...],
) -> None:
    """Render a module-scoped multi-turn AI studio."""
    sessions, active = _studio_sessions(store, module)
    unified = _unified_conversation()
    if unified:
        chat_col = st.container()
    else:
        history_col, chat_col = st.columns([1, 3], gap="small")
        with history_col:
            _studio_history(store, module)
    if st.session_state.pop(f"studio_restore_{module}", False):
        latest = next((str(message.get("content", "")) for message in reversed(active["messages"])
                       if message.get("role") == "assistant" and not message.get("error")), "")
        st.session_state[draft_key] = latest
        st.session_state[editor_key] = latest

    with chat_col:
        st.markdown(
            f"<div class='wb-studio-title'><strong>{html.escape(str(active.get('title') or '新对话'))}</strong>"
            f"<span>已启用 {len(skills)} 个 Skill</span></div>",
            unsafe_allow_html=True,
        )
        with st.container(height=460 if unified else 320, border=True, key=f"studio_chat_messages_{module}"):
            messages = active.setdefault("messages", [])
            if not messages:
                st.markdown(
                    f"<div class='wb-studio-welcome'><div><strong>{html.escape(_label(module))} AI 工作台</strong>"
                    "</div></div>",
                    unsafe_allow_html=True,
                )
            for index, message in enumerate(messages):
                role = "user" if message.get("role") == "user" else "assistant"
                with st.chat_message(role):
                    st.markdown(str(message.get("content", "")))
                    if role == "assistant" and str(message.get("content", "")).strip():
                        if st.button(
                            "设为当前草稿",
                            key=f"studio_use_draft_{module}_{active['id']}_{index}",
                        ):
                            content = str(message["content"])
                            st.session_state[draft_key] = content
                            st.session_state[editor_key] = content
                            st.toast("已放入当前草稿")

        action_columns = st.columns(len(quick_prompts))
        quick_prompt = ""
        for column, (label, prompt_text) in zip(action_columns, quick_prompts):
            if column.button(
                label,
                key=f"studio_quick_{module}_{active['id']}_{label}",
                use_container_width=True,
                disabled=not skills,
            ):
                quick_prompt = prompt_text

        prompt = st.chat_input(
            placeholder,
            key=f"studio_input_{module}",
            disabled=not skills,
        )
        submitted_prompt = quick_prompt or prompt
        if not skills:
            st.info("请先为当前模块选择至少一个 Skill。")
        if submitted_prompt:
            previous_messages = list(active["messages"])
            active["messages"].append({"role": "user", "content": submitted_prompt})
            if active.get("title") == "新对话":
                active["title"] = submitted_prompt.strip().replace("\n", " ")[:18] or "新对话"
            save_session(store, active)
            failed = False
            with st.spinner("AI 正在结合 Skill 和项目内容思考..."):
                try:
                    answer = _generate(
                        build_studio_prompt(
                            module,
                            submitted_prompt,
                            previous_messages,
                            metadata,
                        ),
                        skills,
                        context_text[:18000] + "\n\n当前人工编辑草稿：\n" + str(st.session_state.get(editor_key, ""))[:12000],
                    )
                    st.session_state[draft_key] = answer
                    st.session_state[editor_key] = answer
                except Exception as exc:
                    failed = True
                    answer = f"本次生成失败：{exc}\n\n请检查 API 设置后重试，当前对话已经保存。"
            active["messages"].append({"role": "assistant", "content": answer, "error": failed})
            save_session(store, active)
            st.rerun()


def _distill_context(
    analysis_files: list[Path],
    chapter_files: list[Path],
    review_files: list[Path],
    tracker_files: list[Path],
) -> str:
    """Build a bounded evidence package for conversational distillation."""
    sections: list[str] = []
    for label, paths, per_file_limit, max_files in (
        ("拆解结果", analysis_files, 7000, 2),
        ("章节正文", list(reversed(chapter_files)), 6000, 3),
        ("审查记录", review_files, 3500, 3),
        ("设定追踪", tracker_files, 3500, 4),
    ):
        items: list[str] = []
        for path in paths[:max_files]:
            try:
                text = path.read_text(encoding="utf-8", errors="replace")
            except OSError:
                continue
            items.append(f"### {path.name}\n{text[:per_file_limit]}")
        if items:
            sections.append(f"## {label}\n" + "\n\n".join(items))
    return "\n\n".join(sections)[:30000]


def _creation(store: WorkspaceStore) -> None:
    if not _unified_conversation():
        st.subheader("分步创作")
    skills = _skill_picker("creation")
    targets = {
        "世界观": ("world", "世界背景、力量体系、势力关系和不可破坏的核心限制"),
        "人物关系": ("characters", "人物目标、秘密、关系张力、成长弧与阶段变化"),
        "黄金三章": ("golden", "开篇钩子、连续冲突、首个高潮与每章结尾悬念"),
    }
    with st.expander("创作设置", expanded=False):
        project = st.text_input(
            "项目名称",
            value=st.session_state.get("project_name", "我的新小说"),
            key="creation_project_name",
        )
        st.session_state["project_name"] = project
        target_title = st.radio(
            "当前创作任务",
            list(targets),
            horizontal=True,
            key="creation_studio_target",
        )
        use_context = st.checkbox(
            "使用已保存的项目设定和创作结果",
            value=True,
            key="creation_use_context",
        )

    target_key, target_hint = targets[target_title]
    safe_project = re.sub(r"[^\w\u4e00-\u9fff-]+", "-", project).strip("-") or "我的新小说"
    output_name = f"{safe_project}_{target_key}.md"
    output_path = store.root / "创作结果" / output_name
    saved = output_path.read_text(encoding="utf-8", errors="replace") if output_path.exists() else ""
    draft_key = f"studio_latest_creation_{target_key}"
    editor_key = f"studio_editor_creation_{target_key}"
    if editor_key not in st.session_state:
        st.session_state[editor_key] = st.session_state.get(draft_key, saved)

    context_text = _project_context(store) if use_context else saved
    _studio_chat(
        store,
        "creation",
        skills,
        {
            "项目": project,
            "当前任务": target_title,
            "任务重点": target_hint,
            "已有草稿字符数": len(saved),
        },
        context_text,
        draft_key,
        editor_key,
        "和 AI 讨论世界观、人物或黄金三章...",
        (
            ("提出方案", f"请先为“{target_title}”提出三种可选方向，并说明各自取舍。"),
            ("生成完整稿", f"请根据当前讨论生成一份完整的“{target_title}”草稿。"),
            ("检查漏洞", f"请审查当前“{target_title}”的矛盾、遗漏和后续写作风险，并给出修复稿。"),
        ),
    )

    with st.expander(f"{target_title}草稿与保存", expanded=bool(st.session_state.get(editor_key)) and not _unified_conversation()):
        draft = st.text_area(
            f"{target_title}内容",
            height=420,
            key=editor_key,
            label_visibility="collapsed",
        )
        save_col, open_col = st.columns(2)
        if save_col.button("保存当前成果", type="primary", key=f"save_creation_{target_key}", use_container_width=True):
            path = store.save_text("创作结果", output_name, draft)
            st.session_state["open_file"] = str(path.relative_to(store.root)).replace("\\", "/")
            st.success(f"{target_title}已独立保存")
        if open_col.button("在文件区查看", key=f"open_creation_{target_key}", use_container_width=True, disabled=not output_path.exists()):
            st.session_state["open_file"] = f"创作结果/{output_name}"
            st.rerun()


def _tracking(store: WorkspaceStore) -> None:
    st.markdown("<div class='wb-panel'><h3>设定追踪</h3><div class='wb-muted'>持续维护人物、伏笔、主线和连载状态，避免长篇写作越写越散。</div></div>", unsafe_allow_html=True)
    files = (("人物档案", "characters_tracker.md", "人物姓名、目标、秘密、关系变化和当前状态"), ("伏笔清单", "foreshadowing_tracker.md", "伏笔、首次出现章节、预计回收章节、回收状态"), ("主线与支线", "plot_tracker.md", "主线目标、阶段冲突、支线任务和交汇点"), ("连载状态", "serial_status.md", "当前卷、当前章、下一章目标、读者反馈和待修问题"))
    for title, filename, hint in files:
        path = store.root / "创作结果" / filename
        current = path.read_text(encoding="utf-8", errors="replace") if path.exists() else ""
        with st.container(border=True):
            st.markdown(f"### {title}")
            st.caption(hint)
            value = st.text_area(title, value=current, height=150, key=f"tracker_{filename}")
            if st.button(f"保存{title}", key=f"save_tracker_{filename}"):
                store.save_text("创作结果", filename, value)
                st.success(f"{title}已保存")

def _chapters(store: WorkspaceStore) -> None:
    if not _unified_conversation():
        st.subheader("章节写作")
    skills = _skill_picker("chapters")
    chapter_files = sorted((store.root / "章节内容").glob("第*.md"))
    options = ["新建章节"] + [path.name for path in chapter_files]
    index_path = store.root / "章节内容" / "章节索引.json"
    try:
        chapter_index = json.loads(index_path.read_text(encoding="utf-8")) if index_path.exists() else {}
    except (OSError, ValueError, TypeError):
        chapter_index = {}

    with st.expander("本章设置", expanded=False):
        selected = st.selectbox("选择章节", options, key="chapter_selected")
        existing_meta = chapter_index.get(selected, {}) if selected != "新建章节" else {}
        if selected == "新建章节":
            default_number = max(1, len(chapter_files) + 1)
            default_title = "新的章节"
            existing_content = ""
        else:
            default_number = int(re.search(r"\d+", selected).group()) if re.search(r"\d+", selected) else 1
            default_title = str(existing_meta.get("title") or selected.rsplit("_", 1)[-1].removesuffix(".md"))
            existing_content = store.read_file(f"章节内容/{selected}")
        number = st.number_input(
            "章节编号",
            min_value=1,
            value=default_number,
            step=1,
            key=f"chapter_number_{selected}",
        )
        title = st.text_input(
            "章节标题",
            value=default_title,
            key=f"chapter_title_{selected}",
        )
        outline = st.text_area(
            "本章目标 / 冲突 / 结尾悬念",
            value=str(existing_meta.get("outline", "")),
            height=90,
            key=f"chapter_outline_{selected}",
        )
        statuses = ["草稿", "待审查", "已确认", "已发布"]
        saved_status = str(existing_meta.get("status", "草稿"))
        status = st.selectbox(
            "章节状态",
            statuses,
            index=statuses.index(saved_status) if saved_status in statuses else 0,
            key=f"chapter_status_{selected}",
        )
        use_context = st.checkbox(
            "使用项目设定、创作结果和最近章节",
            value=True,
            key="chapter_use_context",
        )

    draft_key = f"studio_latest_chapters_{selected}"
    editor_key = f"studio_editor_chapters_{selected}"
    if editor_key not in st.session_state:
        st.session_state[editor_key] = st.session_state.get(draft_key, existing_content)
    chapter_context_parts: list[str] = []
    if use_context:
        chapter_context_parts.append(_project_context(store))
        for previous in chapter_files[-3:]:
            if previous.name != selected:
                chapter_context_parts.append(
                    f"最近章节 {previous.name}：\n"
                    f"{previous.read_text(encoding='utf-8', errors='replace')[:6000]}"
                )
    if existing_content:
        chapter_context_parts.append(f"当前章节已有正文：\n{existing_content[:12000]}")

    _studio_chat(
        store,
        "chapters",
        skills,
        {
            "章节编号": int(number),
            "章节标题": title,
            "章节状态": status,
            "本章目标": outline,
            "已有正文字数": len(existing_content),
        },
        "\n\n".join(chapter_context_parts)[:30000],
        draft_key,
        editor_key,
        "和 AI 讨论本章怎么写...",
        (
            ("规划本章", "请根据当前设定规划本章的场景顺序、冲突升级和结尾悬念。"),
            ("生成正文", "请按照当前本章设置和讨论，直接生成可用的完整章节正文。"),
            ("修改润色", "请检查当前章节草稿的节奏、人物一致性和网文阅读感，并给出修改后的完整正文。"),
        ),
    )

    with st.expander("章节草稿与保存", expanded=bool(st.session_state.get(editor_key)) and not _unified_conversation()):
        content = st.text_area(
            "章节正文",
            height=500,
            key=editor_key,
            label_visibility="collapsed",
        )
        save_col, open_col = st.columns(2)
        if save_col.button("保存章节", type="primary", key="save_chapter", use_container_width=True):
            safe_title = re.sub(r"[^\w\u4e00-\u9fff-]+", "-", title).strip("-") or "未命名"
            filename = selected if selected != "新建章节" else f"第{int(number):03d}章_{safe_title}.md"
            path = store.save_text("章节内容", filename, content)
            chapter_index[filename] = {
                "number": int(number),
                "title": title,
                "status": status,
                "outline": outline,
                "updated_at": store.timestamp(),
            }
            store.save_json("章节内容", "章节索引.json", chapter_index)
            st.session_state["open_file"] = str(path.relative_to(store.root)).replace("\\", "/")
            st.success(f"已保存：{path.name}")
        if open_col.button(
            "在文件区查看",
            key="open_saved_chapter",
            use_container_width=True,
            disabled=selected == "新建章节",
        ):
            st.session_state["open_file"] = f"章节内容/{selected}"
            st.rerun()
        st.caption(f"当前草稿字符数：{len(content)} · 已保存章节：{len(chapter_files)}")
    st.caption(f"正文字符数：{len(content)} · 已保存章节：{len(chapter_files)}")


def _review(store: WorkspaceStore) -> None:
    st.markdown("<div class='wb-panel'><h3>质量审查</h3><div class='wb-muted'>本地审查不上传正文，可作为 Skill 进化的客观证据。</div></div>", unsafe_allow_html=True)
    chapter_files = sorted((store.root / "章节内容").glob("第*.md"))
    if not chapter_files:
        st.info("请先在“章节写作”中保存章节")
        return
    chosen = st.selectbox("待审查章节", [path.name for path in chapter_files], key="review_chapter")
    content = store.read_file(f"章节内容/{chosen}")
    if st.button("运行本地质量审查", type="primary", key="run_quality_review"):
        report = review_chapter(content, chosen)
        report["source_file"] = chosen
        report["created_at"] = store.timestamp()
        path = store.save_json("迭代记录", f"review_{store.new_id()}.json", report)
        st.session_state["latest_review"] = report
        st.success(f"审查完成，已保存：{path.name}")
    report = st.session_state.get("latest_review")
    if report and report.get("source_file") == chosen:
        st.metric("本章评分", f"{report['score']}/100")
        st.json(report["metrics"])
        for issue in report["issues"]:
            st.warning(issue["message"] if issue["level"] == "warn" else issue["message"])
        if st.button("将审查结果作为 Skill 进化证据", key="use_review_as_evidence"):
            st.session_state["evolution_evidence"] = json.dumps(report, ensure_ascii=False, indent=2)
            st.session_state["module"] = "feedback"
            st.rerun()

def _feedback(store: WorkspaceStore) -> None:
    st.markdown("<div class='wb-panel'><h3>迭代引擎</h3><div class='wb-muted'>上传、融合、生成或编辑 Skill，并保留每次迭代版本。</div></div>", unsafe_allow_html=True)
    skills = _skill_picker("feedback")
    upload = st.file_uploader("上传新的 Skill", type=["zip"], key="feedback_upload2")
    if upload and st.session_state.get("feedback_upload2_name") != upload.name: import_package(upload.name, upload.getvalue()); st.session_state["feedback_upload2_name"] = upload.name; st.success("Skill 已加入仓库")
    draft = st.text_area("规则草案 / 编辑区", value=st.session_state.get("feedback_draft2", ""), height=220)
    c1, c2 = st.columns(2)
    if c1.button("保存迭代版本", type="primary", disabled=not draft.strip()): store.save_text("迭代记录", f"iteration_{store.new_id()}.md", draft); st.success("迭代版本已保存")
    if c2.button("生成规则草案", disabled=not skills):
        try: st.session_state["feedback_draft2"] = _generate("根据当前 Skill 生成可编辑规则改进草案。", skills); st.rerun()
        except Exception as exc: st.error(f"生成失败：{exc}")
    st.markdown("### Skill 自我进化")
    st.caption("AI 先提出改进提案，只有你确认后才会写入新 Skill 版本。")
    evolution_goal = st.text_area("本轮进化目标", placeholder="例如：减少开篇铺垫，提升前三章悬念密度", key="evolution_goal")
    evidence = st.text_area("依据与问题证据", value=st.session_state.get("evolution_evidence", ""), placeholder="粘贴读者反馈、章节表现或审查发现", key="evolution_evidence")
    if st.button("生成进化提案", key="generate_evolution_proposal", disabled=not (skills and evolution_goal.strip())):
        try:
            prompt = f"请基于当前 Skill，提出一个可执行的 Skill 改进提案。目标：{evolution_goal}\n证据：{evidence}\n必须输出：问题、根因、修改规则、验证指标、回滚条件。"
            st.session_state["evolution_draft"] = _generate(prompt, skills)
            st.rerun()
        except Exception as exc:
            st.error(f"提案生成失败：{exc}")
    proposal = st.text_area("待确认提案", value=st.session_state.get("evolution_draft", ""), height=180, key="evolution_draft_editor")
    if proposal.strip():
        c1, c2 = st.columns(2)
        if c1.button("确认并生成新 Skill 版本", type="primary", key="approve_evolution"):
            evolution_id = store.new_id()
            store.save_json("迭代记录", f"evolution_{evolution_id}.json", {"id": evolution_id, "status": "approved", "goal": evolution_goal, "evidence": evidence, "proposal": proposal, "source_skills": skills, "created_at": store.timestamp()})
            store.save_text("Skill", f"generated_skill_{evolution_id}.md", f"# Generated Skill Evolution\n\n## Goal\n{evolution_goal}\n\n## Approved rules\n{proposal}\n")
            st.session_state.pop("evolution_draft", None)
            st.success("已确认，新的 Skill 版本已保存")
        if c2.button("驳回提案", key="reject_evolution"):
            evolution_id = store.new_id()
            store.save_json("迭代记录", f"evolution_{evolution_id}.json", {"id": evolution_id, "status": "rejected", "goal": evolution_goal, "evidence": evidence, "proposal": proposal, "source_skills": skills, "created_at": store.timestamp()})
            st.session_state.pop("evolution_draft", None)
            st.info("提案已驳回，未修改 Skill")
    _job_controls(store, "feedback")


def _chat(store: WorkspaceStore, module: str) -> None:
    st.markdown("<div class='wb-chat'><strong>⌁ 工作台助手</strong><div class='wb-muted'>对话只服务当前模块，消息保存在本地。</div>", unsafe_allow_html=True)
    messages = st.session_state.setdefault(f"chat_messages2_{module}", [])
    use_context = st.checkbox("本次对话使用项目设定", value=False, key=f"chat_context_{module}")
    for msg in messages[-4:]:
        with st.chat_message(msg["role"]):
            st.write(msg["content"])
    prompt = st.chat_input(f"在“{_label(module)}”中输入任务...")
    if prompt:
        messages.append({"role": "user", "content": prompt})
        try: answer = _generate(prompt, st.session_state.get(f"selected_{module}", []), _project_context(store) if use_context else "")
        except Exception as exc: answer = f"还不能调用模型：{exc}\n请先到 API设置 填写接口，或继续使用本地保存功能。"
        messages.append({"role": "assistant", "content": answer}); store.save_json("迭代记录", f"chat_{module}.json", {"module": module, "messages": messages}); st.rerun()
    st.markdown("</div>", unsafe_allow_html=True)


def _home(store: WorkspaceStore) -> None:
    st.markdown("<div class='wb-panel'><h3>今天从哪里开始？</h3><div class='wb-muted'>素材、规则和创作结果都放在同一个本地工作区里，按任务继续推进。</div></div>", unsafe_allow_html=True)
    profile = store.project_profile()
    with st.expander("小说项目设定", expanded=True):
        profile["name"] = st.text_input("书名", value=profile.get("name", "我的新小说"), key="profile_name")
        profile["genre"] = st.text_input("题材 / 目标平台", value=profile.get("genre", ""), key="profile_genre")
        profile["premise"] = st.text_area("一句话故事核", value=profile.get("premise", ""), height=70, key="profile_premise")
        profile["audience"] = st.text_input("目标读者与爽点", value=profile.get("audience", ""), key="profile_audience")
        if st.button("保存项目设定", key="save_project_profile"):
            store.save_project_profile(profile)
            st.success("项目设定已保存")
    jobs = JobManager(store).list_jobs(); counts = [("素材", len(list((store.root / "小说资料").glob("*"))), "已导入"), ("拆解", len(list((store.root / "拆解结果").glob("*"))), "份结果"), ("创作", len(list((store.root / "创作结果").glob("*"))), "份产物"), ("任务", len(jobs), "条记录")]
    cols = st.columns(4)
    for col, (label, value, suffix) in zip(cols, counts):
        with col: st.markdown(f"<div class='wb-stat'><div class='wb-muted'>{label} · {suffix}</div><div class='wb-stat-value'>{value}</div></div>", unsafe_allow_html=True)
    st.markdown("<div class='wb-panel'><h3>最近任务</h3>", unsafe_allow_html=True)
    for job in jobs[:6]: st.write(f"{_label(job.module)} · {job.title} · {job.status}")
    if not jobs: st.caption("还没有任务记录")
    st.markdown("</div>", unsafe_allow_html=True)


def _conversation(store: WorkspaceStore) -> None:
    last_mode = st.session_state.get("conversation_last_mode", "material")
    mode = st.radio(
        "对话模式",
        CONVERSATION_MODES,
        format_func=_label,
        horizontal=True,
        index=CONVERSATION_MODES.index(last_mode) if last_mode in CONVERSATION_MODES else 0,
        key="conversation_mode",
        label_visibility="collapsed",
    )
    st.session_state["conversation_last_mode"] = mode
    renderers = {
        "material": _material,
        "creation": _creation,
        "chapters": _chapters,
        "distill": _distill,
    }
    renderers[mode](store)


def render() -> None:
    store = WorkspaceStore(DEFAULT_WORKSPACE_ID)
    _css(); module = _sidebar(store)
    _header(module)
    main_col, files_col = st.columns(WORKBENCH_COLUMNS, gap="medium")
    with main_col:
        if module == "home": _home(store)
        elif module == "conversation": _conversation(store)
        elif module == "material": _material(store)
        elif module == "distill": _distill(store)
        elif module == "skill": _skill(store)
        elif module == "creation": _creation(store)
        elif module == "chapters": _chapters(store)
        elif module == "tracking": _tracking(store)
        elif module == "feedback": _feedback(store)
        elif module == "review": _review(store)
        elif module == "settings":
            from ui.pages.settings_page import render as settings_render
            settings_render()
        else: _home(store)
        if module not in ("settings", "conversation", *STUDIO_MODULES): _chat(store, module)
    with files_col:
        _file_workspace(store)

















